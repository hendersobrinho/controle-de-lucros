"""Tela de distribuição trimestral e o reflexo dela na aba anual."""
import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import sqlite3

import pytest
from PySide6.QtWidgets import QApplication, QMessageBox

from controle_lucros import db, repositories as repo
from controle_lucros.models import Empresa, Socio
from controle_lucros.ui import distribuicao_trimestral_view as mod
from controle_lucros.ui.distribuicao_anual_view import DistribuicaoAnualView
from controle_lucros.ui.distribuicao_trimestral_view import (
    COL_IRRF,
    COL_PRO_LABORE,
    COL_VALOR,
    DistribuicaoTrimestralView,
)


@pytest.fixture(scope="module", autouse=True)
def app():
    aplicativo = QApplication.instance() or QApplication([])
    yield aplicativo


@pytest.fixture()
def conn():
    connection = sqlite3.connect(":memory:")
    connection.row_factory = sqlite3.Row
    db.init_schema(connection)
    yield connection
    connection.close()


@pytest.fixture()
def cenario(conn):
    empresa_id = repo.salvar_empresa(conn, Empresa(None, "001", "ACME LTDA", "", 100000, 1000))
    fulano = repo.salvar_socio(conn, Socio(None, "Fulano de Tal", "111.111.111-11"))
    beltrano = repo.salvar_socio(conn, Socio(None, "Beltrano da Silva", "222.222.222-22"))
    repo.associar_socio_a_empresa(conn, empresa_id, fulano, 60.0, 600, "2020-01-01")
    repo.associar_socio_a_empresa(conn, empresa_id, beltrano, 40.0, 400, "2020-01-01")
    return {"empresa": empresa_id, "fulano": fulano, "beltrano": beltrano}


def _view(conn, trimestre: int = 1) -> DistribuicaoTrimestralView:
    view = DistribuicaoTrimestralView(conn)
    view.ano.setValue(2025)
    view.trimestre.setCurrentIndex(trimestre - 1)
    return view


def _linha_de(view, socio_id: int) -> int:
    return next(i for i, l in enumerate(view._linhas) if l["socio_id"] == socio_id)


def _lancar(view, socio_id: int, valor: float, pro_labore: float = 0.0, irrf: float = 0.0) -> None:
    view._iniciar_edicao()
    row = _linha_de(view, socio_id)
    view.tabela.cellWidget(row, COL_VALOR).setValue(valor)
    view.tabela.cellWidget(row, COL_PRO_LABORE).setValue(pro_labore)
    view.tabela.cellWidget(row, COL_IRRF).setValue(irrf)
    view._salvar_edicao()


@pytest.fixture(autouse=True)
def sem_popups(monkeypatch):
    monkeypatch.setattr(mod.QMessageBox, "information", lambda *a, **k: None)


def test_abre_no_trimestre_corrente(conn, cenario):
    view = DistribuicaoTrimestralView(conn)
    assert view.trimestre.currentData() == (mod.dt.date.today().month - 1) // 3 + 1


def test_lista_os_socios_da_empresa_no_trimestre(conn, cenario):
    view = _view(conn)
    assert view.tabela.rowCount() == 2
    assert {l["socio_id"] for l in view._linhas} == {cenario["fulano"], cenario["beltrano"]}
    assert "Nenhum trimestre lançado" in view.progresso.text()


def test_lancar_trimestre_grava_e_reflete_no_anual(conn, cenario):
    view = _view(conn, trimestre=1)
    _lancar(view, cenario["fulano"], 10000.0, pro_labore=6000.0, irrf=138.0)

    assert view._linhas[_linha_de(view, cenario["fulano"])]["valor_distribuido"] == 10000.0
    anual = next(
        d for d in repo.listar_distribuicoes(conn, cenario["empresa"], 2025)
        if d.socio_id == cenario["fulano"]
    )
    assert (anual.valor_distribuido, anual.pro_labore, anual.irrf) == (10000.0, 6000.0, 138.0)


def test_acumulado_soma_os_trimestres_anteriores(conn, cenario):
    _lancar(_view(conn, trimestre=1), cenario["fulano"], 10000.0)
    view = _view(conn, trimestre=2)
    _lancar(view, cenario["fulano"], 20000.0)

    linha = view._linhas[_linha_de(view, cenario["fulano"])]
    assert linha["valor_distribuido"] == 20000.0
    assert linha["acumulado_valor"] == 30000.0
    assert "1º, 2º" in view.progresso.text()
    assert "30.000,00" in view.progresso.text()


def test_acumulado_do_ano_conta_socio_que_saiu_antes(conn, cenario):
    """O total do ano não pode encolher só porque o sócio não aparece mais na
    lista do trimestre em que se está olhando."""
    _lancar(_view(conn, trimestre=1), cenario["beltrano"], 8000.0)
    vinculo = next(
        v for v in repo.listar_vinculos_socio(conn, cenario["beltrano"]) if v.data_saida is None
    )
    repo.encerrar_vinculo_registrando_alteracao(conn, vinculo, "2025-02-15", "Saída")

    view = _view(conn, trimestre=4)
    _lancar(view, cenario["fulano"], 2000.0)
    assert cenario["beltrano"] not in {l["socio_id"] for l in view._linhas}
    assert "10.000,00" in view.progresso.text()  # 8.000 do beltrano + 2.000 do fulano


def test_limpar_trimestre_apaga_lancamentos_e_recalcula_o_anual(conn, cenario, monkeypatch):
    _lancar(_view(conn, trimestre=1), cenario["fulano"], 10000.0)
    view = _view(conn, trimestre=2)
    _lancar(view, cenario["fulano"], 20000.0)

    monkeypatch.setattr(mod.QMessageBox, "question", lambda *a, **k: QMessageBox.Yes)
    view._limpar_trimestre()

    assert repo.trimestres_lancados(conn, cenario["empresa"], 2025) == [1]
    anual = next(
        d for d in repo.listar_distribuicoes(conn, cenario["empresa"], 2025)
        if d.socio_id == cenario["fulano"]
    )
    assert anual.valor_distribuido == 10000.0


def test_periodo_trancado_bloqueia_a_tela(conn, cenario):
    repo.fechar_periodo(conn, cenario["empresa"], 2025)
    view = _view(conn)
    assert view.aviso_trancado.isVisible() or view._periodo_fechado
    assert not view.btn_editar.isEnabled()
    assert not view.btn_limpar.isEnabled()


def test_limpar_fica_desabilitado_sem_lancamento(conn, cenario):
    view = _view(conn)
    assert view.btn_editar.isEnabled()
    assert not view.btn_limpar.isEnabled()


def test_salvar_sem_mudar_nada_nao_cria_lancamento(conn, cenario):
    view = _view(conn)
    view._iniciar_edicao()
    view._salvar_edicao()
    assert repo.listar_distribuicoes_trimestrais(conn, cenario["empresa"], 2025) == []


# --------------------------------------------------- reflexo na aba anual --


def test_aba_anual_mostra_origem_trimestres(conn, cenario):
    _lancar(_view(conn, trimestre=1), cenario["fulano"], 10000.0)

    anual = DistribuicaoAnualView(conn)
    anual.ano.setValue(2025)
    row = next(i for i, l in enumerate(anual._linhas) if l["socio_id"] == cenario["fulano"])
    from controle_lucros.ui.distribuicao_anual_view import COL_ORIGEM

    assert anual.tabela.item(row, COL_ORIGEM).text() == "trimestres"
    assert "acumulado dos trimestres lançados (1º)" in anual.resumo.text()


def test_aba_anual_avisa_quando_o_valor_foi_editado_a_mao(conn, cenario):
    _lancar(_view(conn, trimestre=1), cenario["fulano"], 10000.0)
    repo.salvar_distribuicao(conn, cenario["empresa"], 2025, cenario["fulano"], 12345.0)

    anual = DistribuicaoAnualView(conn)
    anual.ano.setValue(2025)
    row = next(i for i, l in enumerate(anual._linhas) if l["socio_id"] == cenario["fulano"])
    from controle_lucros.ui.distribuicao_anual_view import COL_ORIGEM

    texto = anual.tabela.item(row, COL_ORIGEM).text()
    assert texto.startswith("editado à mão")
    assert "10.000,00" in texto  # mostra o somatório de que ele divergiu


def test_aba_anual_sem_trimestres_nao_mostra_origem(conn, cenario):
    repo.salvar_distribuicao(conn, cenario["empresa"], 2025, cenario["fulano"], 50000.0)
    anual = DistribuicaoAnualView(conn)
    anual.ano.setValue(2025)
    row = next(i for i, l in enumerate(anual._linhas) if l["socio_id"] == cenario["fulano"])
    from controle_lucros.ui.distribuicao_anual_view import COL_ORIGEM

    assert anual.tabela.item(row, COL_ORIGEM).text() == "—"
    assert "acumulado dos trimestres" not in anual.resumo.text()


def test_edicao_em_linha_da_aba_anual_continua_nas_colunas_certas(conn, cenario):
    """A coluna "Origem" entrou no meio da tabela; os widgets de edição são
    posicionados por índice, então precisam continuar caindo na célula certa."""
    from controle_lucros.ui.distribuicao_anual_view import (
        COL_COTAS,
        COL_IRRF as ANUAL_IRRF,
        COL_ORIGEM as ANUAL_ORIGEM,
        COL_PERCENTUAL,
        COL_PRO_LABORE as ANUAL_PRO_LABORE,
        COL_VALOR as ANUAL_VALOR,
    )

    repo.salvar_distribuicao(conn, cenario["empresa"], 2025, cenario["fulano"], 50000.0, 3000.0, 90.0)
    anual = DistribuicaoAnualView(conn)
    anual.ano.setValue(2025)
    anual._iniciar_edicao()
    row = next(i for i, l in enumerate(anual._linhas) if l["socio_id"] == cenario["fulano"])

    assert anual.tabela.cellWidget(row, ANUAL_VALOR).value() == 50000.0
    assert anual.tabela.cellWidget(row, ANUAL_PRO_LABORE).value() == 3000.0
    assert anual.tabela.cellWidget(row, ANUAL_IRRF).value() == 90.0
    assert anual.tabela.cellWidget(row, COL_PERCENTUAL).value() == 60.0
    assert anual.tabela.cellWidget(row, COL_COTAS).value() == 600.0
    # A coluna de origem não é editável: não pode ganhar widget.
    assert anual.tabela.cellWidget(row, ANUAL_ORIGEM) is None


# ------------------------------------------------------- planilha do trimestre --


@pytest.fixture()
def sem_dialogos(monkeypatch, tmp_path):
    """Confirmações e seletor de arquivo são modais; sem interceptar, o teste
    trava esperando um clique."""
    monkeypatch.setattr(mod.QMessageBox, "question", lambda *a, **k: QMessageBox.Yes)
    monkeypatch.setattr(mod.QMessageBox, "warning", lambda *a, **k: None)
    return tmp_path


def _exportar(view, caminho, monkeypatch):
    monkeypatch.setattr(mod.QFileDialog, "getSaveFileName", lambda *a, **k: (str(caminho), ""))
    view._exportar_modelo()


def _importar(view, caminho, monkeypatch):
    monkeypatch.setattr(mod.QFileDialog, "getOpenFileName", lambda *a, **k: (str(caminho), ""))
    view._importar_planilha()


def test_exporta_modelo_com_os_socios_do_trimestre(conn, cenario, sem_dialogos, monkeypatch):
    from controle_lucros.planilha import importar_distribuicao

    view = _view(conn, trimestre=1)
    _lancar(view, cenario["fulano"], 10000.0, pro_labore=6000.0, irrf=138.0)

    caminho = sem_dialogos / "t1.xlsx"
    _exportar(view, caminho, monkeypatch)

    linhas = importar_distribuicao(caminho)
    assert {l["nome"] for l in linhas} == {"Fulano de Tal", "Beltrano da Silva"}
    fulano = next(l for l in linhas if l["nome"] == "Fulano de Tal")
    # Vem preenchido com o que já foi lançado: serve de conferência também.
    assert (fulano["valor_distribuido"], fulano["pro_labore"], fulano["irrf"]) == (10000.0, 6000.0, 138.0)


def test_importar_lanca_no_trimestre_da_tela_e_reflete_no_anual(conn, cenario, sem_dialogos, monkeypatch):
    view = _view(conn, trimestre=2)
    caminho = sem_dialogos / "t2.xlsx"
    _exportar(view, caminho, monkeypatch)

    import openpyxl

    wb = openpyxl.load_workbook(caminho)
    aba = wb["Distribuição"]
    for linha in aba.iter_rows(min_row=2):
        if linha[1].value == "Fulano de Tal":
            linha[2].value, linha[3].value, linha[4].value = 20000, 9000, 200
    wb.save(caminho)

    _importar(view, caminho, monkeypatch)

    (lancamento,) = [
        d for d in repo.listar_distribuicoes_trimestrais(conn, cenario["empresa"], 2025, trimestre=2)
        if d.socio_id == cenario["fulano"]
    ]
    assert (lancamento.valor_distribuido, lancamento.pro_labore, lancamento.irrf) == (20000.0, 9000.0, 200.0)

    anual = next(
        d for d in repo.listar_distribuicoes(conn, cenario["empresa"], 2025)
        if d.socio_id == cenario["fulano"]
    )
    assert anual.valor_distribuido == 20000.0


def test_importar_nao_toca_nos_outros_trimestres(conn, cenario, sem_dialogos, monkeypatch):
    _lancar(_view(conn, trimestre=1), cenario["fulano"], 10000.0)

    view = _view(conn, trimestre=3)
    caminho = sem_dialogos / "t3.xlsx"
    _exportar(view, caminho, monkeypatch)
    import openpyxl

    wb = openpyxl.load_workbook(caminho)
    for linha in wb["Distribuição"].iter_rows(min_row=2):
        if linha[1].value == "Fulano de Tal":
            linha[2].value = 5000
    wb.save(caminho)
    _importar(view, caminho, monkeypatch)

    assert repo.trimestres_lancados(conn, cenario["empresa"], 2025) == [1, 3]
    anual = next(
        d for d in repo.listar_distribuicoes(conn, cenario["empresa"], 2025)
        if d.socio_id == cenario["fulano"]
    )
    assert anual.valor_distribuido == 15000.0  # 10.000 do 1º + 5.000 do 3º


def test_trimestre_em_branco_importa_sem_perguntar_nada(conn, cenario, tmp_path, monkeypatch):
    """O trimestre já está selecionado no topo da tela — repetir isso numa
    caixa de confirmação só cobraria um clique a mais. Sem lançamento pra
    substituir, não há o que perguntar."""
    perguntas = []
    monkeypatch.setattr(
        mod.QMessageBox, "question", lambda *a, **k: (perguntas.append(a[2]), QMessageBox.Yes)[1]
    )
    monkeypatch.setattr(mod.QMessageBox, "information", lambda *a, **k: None)

    view = _view(conn, trimestre=1)
    caminho = tmp_path / "t.xlsx"
    _exportar(view, caminho, monkeypatch)

    import openpyxl

    wb = openpyxl.load_workbook(caminho)
    for linha in wb["Distribuição"].iter_rows(min_row=2):
        linha[2].value = 999
    wb.save(caminho)

    _importar(view, caminho, monkeypatch)

    assert perguntas == []
    assert len(repo.listar_distribuicoes_trimestrais(conn, cenario["empresa"], 2025, trimestre=1)) == 2


def test_substituir_lancamento_existente_pede_confirmacao(conn, cenario, tmp_path, monkeypatch):
    """Aqui há o que perder: o valor antigo não volta."""
    monkeypatch.setattr(mod.QMessageBox, "information", lambda *a, **k: None)
    view = _view(conn, trimestre=1)
    _lancar(view, cenario["fulano"], 10000.0)

    caminho = tmp_path / "t.xlsx"
    _exportar(view, caminho, monkeypatch)
    import openpyxl

    wb = openpyxl.load_workbook(caminho)
    for linha in wb["Distribuição"].iter_rows(min_row=2):
        if linha[1].value == "Fulano de Tal":
            linha[2].value = 777
    wb.save(caminho)

    perguntas = []
    monkeypatch.setattr(
        mod.QMessageBox, "question", lambda *a, **k: (perguntas.append(a[2]), QMessageBox.No)[1]
    )
    _importar(view, caminho, monkeypatch)

    assert perguntas and "já têm valor lançado" in perguntas[0]
    (lancamento,) = [
        d for d in repo.listar_distribuicoes_trimestrais(conn, cenario["empresa"], 2025, trimestre=1)
        if d.socio_id == cenario["fulano"]
    ]
    assert lancamento.valor_distribuido == 10000.0  # recusou: nada foi substituído


def test_importar_com_periodo_trancado_avisa_e_nao_lanca(conn, cenario, tmp_path, monkeypatch):
    view = _view(conn, trimestre=1)
    caminho = tmp_path / "t.xlsx"
    monkeypatch.setattr(mod.QMessageBox, "question", lambda *a, **k: QMessageBox.Yes)
    _exportar(view, caminho, monkeypatch)

    repo.fechar_periodo(conn, cenario["empresa"], 2025)
    avisos = []
    monkeypatch.setattr(mod.QMessageBox, "warning", lambda *a, **k: avisos.append(a[2]))
    _importar(view, caminho, monkeypatch)

    assert avisos and "trancado" in avisos[0]
    assert repo.listar_distribuicoes_trimestrais(conn, cenario["empresa"], 2025) == []


def test_socio_desconhecido_na_planilha_vai_pra_revisao(conn, cenario, sem_dialogos, monkeypatch):
    """Nunca cadastra sócio sozinho — mesma regra da aba anual."""
    from controle_lucros.ui.importacao_distribuicao import associar_linhas

    linhas = [{"cpf": "999.999.999-99", "nome": "Ninguém Conhecido", "valor_distribuido": 1.0,
               "pro_labore": 0.0, "irrf": 0.0}]
    resolvidos, pendencias = associar_linhas(conn, linhas, {cenario["fulano"]}, "no 1º trimestre de 2025")
    assert resolvidos == []
    assert len(pendencias) == 1
    assert "Nenhum sócio cadastrado" in pendencias[0]["aviso"]


def test_socio_fora_do_trimestre_vira_pendencia_com_o_periodo_no_aviso(conn, cenario):
    from controle_lucros.ui.importacao_distribuicao import associar_linhas

    linhas = [{"cpf": "222.222.222-22", "nome": "Beltrano da Silva", "valor_distribuido": 1.0,
               "pro_labore": 0.0, "irrf": 0.0}]
    _, pendencias = associar_linhas(conn, linhas, {cenario["fulano"]}, "no 4º trimestre de 2025")
    assert len(pendencias) == 1
    assert "no 4º trimestre de 2025" in pendencias[0]["aviso"]


def test_botoes_de_planilha_seguem_o_estado_da_tela(conn, cenario):
    view = _view(conn)
    assert view.btn_exportar_modelo.isEnabled()
    assert view.btn_importar.isEnabled()

    repo.fechar_periodo(conn, cenario["empresa"], 2025)
    view = _view(conn)
    # Trancado: dá pra exportar pra conferir, mas não pra importar.
    assert view.btn_exportar_modelo.isEnabled()
    assert not view.btn_importar.isEnabled()
