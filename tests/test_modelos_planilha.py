"""Modelos de planilha do cadastro em massa e a aba de exemplo.

Os três modelos são lidos pelo MESMO importador — colunas identificadas pelo
cabeçalho — então um modelo menor é só um arquivo com menos colunas, não um
formato novo. É isso que estes testes prendem.
"""
import sqlite3

import openpyxl
import pytest

from controle_lucros import db, repositories as repo
from controle_lucros.planilha import (
    COLUNAS_CADASTRO,
    LINHAS_EXEMPLO_CADASTRO,
    MODELOS_CADASTRO,
    exportar_modelo_cadastro,
    exportar_modelo_distribuicao,
    importar_cadastro,
    importar_distribuicao,
    modelo_cadastro,
)


@pytest.fixture()
def conn():
    connection = sqlite3.connect(":memory:")
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON;")
    db.init_schema(connection)
    yield connection
    connection.close()


IDS = [m.id for m in MODELOS_CADASTRO]


# ------------------------------------------------------------- os modelos --


def test_modelos_tem_ids_unicos_e_descricao():
    assert len(IDS) == len(set(IDS))
    for modelo in MODELOS_CADASTRO:
        assert modelo.nome.strip() and modelo.descricao.strip()


def test_toda_coluna_de_modelo_existe_no_modelo_geral():
    """Coluna que não estiver em COLUNAS_CADASTRO não seria reconhecida pelo
    importador — sairia da planilha e entraria como campo vazio."""
    for modelo in MODELOS_CADASTRO:
        assert set(modelo.colunas) <= set(COLUNAS_CADASTRO), modelo.id


def test_modelo_geral_tem_todas_as_colunas():
    assert modelo_cadastro("geral").colunas == COLUNAS_CADASTRO


def test_modelos_menores_sao_de_fato_menores():
    assert len(modelo_cadastro("empresas").colunas) < len(modelo_cadastro("empresas_socios").colunas)
    assert len(modelo_cadastro("empresas_socios").colunas) < len(modelo_cadastro("geral").colunas)


def test_modelo_desconhecido_da_erro_claro():
    with pytest.raises(ValueError, match="desconhecido"):
        modelo_cadastro("nao-existe")


# ------------------------------------------------------- a aba de exemplo --


@pytest.mark.parametrize("id_modelo", IDS)
def test_planilha_sai_com_aba_de_dados_e_aba_de_exemplo(tmp_path, id_modelo):
    caminho = tmp_path / f"{id_modelo}.xlsx"
    exportar_modelo_cadastro(caminho, modelo=id_modelo)
    workbook = openpyxl.load_workbook(caminho)
    assert workbook.sheetnames == ["Cadastro", "Exemplo"]
    # A aba de dados precisa ser a ativa: é a que o importador lê e a que a
    # pessoa vê ao abrir o arquivo.
    assert workbook.active.title == "Cadastro"


@pytest.mark.parametrize("id_modelo", IDS)
def test_o_exemplo_nao_e_importado_junto(tmp_path, id_modelo):
    """O risco de pôr exemplo na planilha é alguém importar sem apagar e criar
    empresa fictícia no banco. Por isso ele fica em aba separada."""
    caminho = tmp_path / f"{id_modelo}.xlsx"
    exportar_modelo_cadastro(caminho, modelo=id_modelo)
    assert importar_cadastro(caminho) == []


@pytest.mark.parametrize("id_modelo", IDS)
def test_o_exemplo_tem_o_mesmo_cabecalho_do_modelo(tmp_path, id_modelo):
    """Exemplo com colunas diferentes das da aba de dados ensinaria errado na
    hora de copiar e colar."""
    modelo = modelo_cadastro(id_modelo)
    caminho = tmp_path / f"{id_modelo}.xlsx"
    exportar_modelo_cadastro(caminho, modelo=id_modelo)
    aba = openpyxl.load_workbook(caminho)["Exemplo"]

    cabecalhos = [
        [c.value for c in linha if c.value is not None]
        for linha in aba.iter_rows()
        if linha and linha[0].value == modelo.colunas[0]
    ]
    assert cabecalhos == [modelo.colunas]


def test_exemplo_mostra_empresa_com_varios_socios_e_socio_em_varias_empresas():
    """São as duas dúvidas que o exemplo existe para responder."""
    por_empresa = {}
    por_socio = {}
    for linha in LINHAS_EXEMPLO_CADASTRO:
        por_empresa.setdefault(linha["empresa_nome"], set()).add(linha["socio_nome"])
        por_socio.setdefault(linha["socio_nome"], set()).add(linha["empresa_nome"])

    assert any(len(socios) > 1 for socios in por_empresa.values())
    assert any(len(empresas) > 1 for empresas in por_socio.values())
    assert any(l["tipo_pessoa"].lower().startswith("j") for l in LINHAS_EXEMPLO_CADASTRO)
    assert any(l["data_saida"] for l in LINHAS_EXEMPLO_CADASTRO)


def _linhas_de_exemplo(caminho, modelo) -> list[list]:
    """As linhas de dados da aba Exemplo: o que vem depois do cabeçalho."""
    aba = openpyxl.load_workbook(caminho)["Exemplo"]
    linhas = [[c.value for c in linha] for linha in aba.iter_rows()]
    inicio = next(i for i, l in enumerate(linhas) if l and l[0] == modelo.colunas[0])
    return [l for l in linhas[inicio + 1:] if l and l[0] is not None]


def test_exemplo_de_so_empresas_traz_uma_linha_por_empresa(tmp_path):
    """Sem coluna de sócio, repetir a empresa não ensina nada e faria parecer
    que o exemplo tem cinco empresas quando tem três."""
    modelo = modelo_cadastro("empresas")
    caminho = tmp_path / "empresas.xlsx"
    exportar_modelo_cadastro(caminho, modelo=modelo)

    nomes = [linha[1] for linha in _linhas_de_exemplo(caminho, modelo)]
    assert nomes == sorted(set(nomes), key=nomes.index)
    assert len(nomes) == len({l["empresa_nome"] for l in LINHAS_EXEMPLO_CADASTRO})


def test_exemplo_do_modelo_geral_repete_a_empresa_por_socio(tmp_path):
    """No modelo completo a repetição é justamente o que precisa ser mostrado."""
    modelo = modelo_cadastro("geral")
    caminho = tmp_path / "geral.xlsx"
    exportar_modelo_cadastro(caminho, modelo=modelo)

    linhas = _linhas_de_exemplo(caminho, modelo)
    assert len(linhas) == len(LINHAS_EXEMPLO_CADASTRO)
    nomes = [linha[1] for linha in linhas]
    assert len(nomes) > len(set(nomes))


# ------------------------------------------------ importar cada modelo --


def _preencher(caminho, linhas: list[list]) -> None:
    workbook = openpyxl.load_workbook(caminho)
    aba = workbook["Cadastro"]
    for linha in linhas:
        aba.append(linha)
    workbook.save(caminho)


def test_modelo_so_empresas_cadastra_sem_exigir_socio(tmp_path, conn):
    caminho = tmp_path / "empresas.xlsx"
    exportar_modelo_cadastro(caminho, modelo="empresas")
    _preencher(caminho, [
        ["010", "ALFA LTDA", "55.555.555/0001-55", 80000, 80000],
        ["011", "BETA LTDA", "", 40000, 4000],
    ])

    linhas = importar_cadastro(caminho)
    assert len(linhas) == 2
    resultado = repo.preparar_importacao_cadastro(conn, linhas)
    assert resultado["pendencias"] == []
    assert all(p["socio_id"] is None for p in resultado["prontas"])

    aplicado = repo.aplicar_importacao_cadastro(conn, resultado["prontas"])
    assert aplicado["empresas_criadas"] == 2
    assert aplicado["vinculos_criados"] == 0
    assert sorted(e.nome for e in repo.listar_empresas(conn)) == ["ALFA LTDA", "BETA LTDA"]


def test_modelo_empresas_socios_cria_empresa_e_vinculo(tmp_path, conn):
    caminho = tmp_path / "es.xlsx"
    exportar_modelo_cadastro(caminho, modelo="empresas_socios")
    _preencher(caminho, [
        ["020", "GAMA LTDA", "", 100000, 1000, "Fulano de Tal", "111.111.111-11", "Física", 60, 600, "01/01/2020"],
        ["020", "GAMA LTDA", "", 100000, 1000, "Beltrano da Silva", "222.222.222-22", "Física", 40, 400, "01/01/2020"],
    ])

    linhas = importar_cadastro(caminho)
    assert len(linhas) == 2
    resultado = repo.preparar_importacao_cadastro(conn, linhas)
    # Sócios ainda não existem: viram pendência, nunca cadastro automático.
    assert len(resultado["pendencias"]) == 2

    for socio_nome, cpf in (("Fulano de Tal", "111.111.111-11"), ("Beltrano da Silva", "222.222.222-22")):
        from controle_lucros.models import Socio

        repo.salvar_socio(conn, Socio(id=None, nome=socio_nome, cpf=cpf))

    resultado = repo.preparar_importacao_cadastro(conn, linhas)
    assert resultado["pendencias"] == []
    aplicado = repo.aplicar_importacao_cadastro(conn, resultado["prontas"])
    assert aplicado["empresas_criadas"] == 1
    assert aplicado["vinculos_criados"] == 2


def test_modelo_enxuto_nao_apaga_o_que_nao_esta_nele(tmp_path, conn):
    """O modelo "Só empresas" não tem coluna de sócio — importar por ele não
    pode encerrar vínculo nem zerar distribuição de quem já está cadastrado."""
    from controle_lucros.models import Empresa, Socio

    empresa_id = repo.salvar_empresa(conn, Empresa(None, "030", "DELTA LTDA", "", 1000, 100))
    socio_id = repo.salvar_socio(conn, Socio(None, "Fulano de Tal", "111.111.111-11"))
    repo.associar_socio_a_empresa(conn, empresa_id, socio_id, 100.0, 100, "2020-01-01")
    repo.salvar_distribuicao(conn, empresa_id, 2025, socio_id, 5000.0)

    caminho = tmp_path / "empresas.xlsx"
    exportar_modelo_cadastro(caminho, modelo="empresas")
    _preencher(caminho, [["030", "DELTA LTDA", "", 1000, 100]])

    resultado = repo.preparar_importacao_cadastro(conn, importar_cadastro(caminho))
    repo.aplicar_importacao_cadastro(conn, resultado["prontas"])

    vinculos = repo.listar_vinculos_empresa(conn, empresa_id)
    assert len(vinculos) == 1 and vinculos[0].data_saida is None
    (distribuicao,) = repo.listar_distribuicoes(conn, empresa_id, 2025)
    assert distribuicao.valor_distribuido == 5000.0


def test_coluna_de_socio_em_branco_continua_sendo_linha_ignorada(tmp_path):
    """Coluna ausente é decisão do modelo; coluna presente e vazia é descuido
    — e continua sendo ignorada, pra não criar empresa por engano."""
    caminho = tmp_path / "geral.xlsx"
    exportar_modelo_cadastro(caminho, modelo="geral")
    _preencher(caminho, [["040", "OMEGA LTDA", "", 1000, 100, "", "", "fisica", 0, 0, ""]])
    assert importar_cadastro(caminho) == []


# ------------------------------------------------ modelo de distribuição --


def test_modelo_de_distribuicao_tambem_tem_aba_de_exemplo(tmp_path):
    caminho = tmp_path / "dist.xlsx"
    exportar_modelo_distribuicao(caminho, [{"cpf": "111.111.111-11", "nome": "Fulano"}])
    workbook = openpyxl.load_workbook(caminho)
    assert workbook.sheetnames == ["Distribuição", "Exemplo"]
    assert workbook.active.title == "Distribuição"

    # O exemplo não pode entrar junto: só o sócio de verdade é importado.
    (linha,) = importar_distribuicao(caminho)
    assert linha["nome"] == "Fulano"
