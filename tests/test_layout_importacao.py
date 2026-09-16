"""Layout de importação: descrever por letra onde cada informação está.

O que se testa aqui é a promessa central do recurso — importar uma planilha
que NÃO segue o modelo do sistema, com as colunas em qualquer ordem, buracos
no meio e o cabeçalho onde a origem quis pôr.
"""
import openpyxl
import pytest

from controle_lucros.layout_importacao import (
    CAMPOS_POR_GRUPO,
    GRUPOS_CAMPOS,
    LayoutImportacao,
    LayoutInvalido,
    exportar_com_layout,
    importar_com_layout,
    indice_para_letra,
    layout_do_modelo,
    letra_para_indice,
    limpar_colunas,
    previa,
    validar,
)
from controle_lucros.planilha import CAMPOS_CADASTRO, modelo_cadastro


def _planilha(tmp_path, linhas, nome="origem.xlsx"):
    caminho = tmp_path / nome
    workbook = openpyxl.Workbook()
    aba = workbook.active
    for linha in linhas:
        aba.append(linha)
    workbook.save(caminho)
    return caminho


def _layout(**colunas):
    linha_inicial = colunas.pop("linha_inicial", 2)
    return LayoutImportacao("Origem", limpar_colunas(colunas), linha_inicial)


# ------------------------------------------------------------------ letras --
@pytest.mark.parametrize(
    "letra,indice", [("A", 0), ("a", 0), ("Z", 25), ("AA", 26), ("AZ", 51), ("BA", 52)]
)
def test_letra_de_coluna_vira_posicao(letra, indice):
    assert letra_para_indice(letra) == indice


@pytest.mark.parametrize("letra", ["", "1", "A1", "ABCD", "Ç", " "])
def test_letra_invalida_reclama(letra):
    with pytest.raises(LayoutInvalido):
        letra_para_indice(letra)


def test_ida_e_volta_da_letra():
    assert [indice_para_letra(i) for i in (0, 25, 26, 701)] == ["A", "Z", "AA", "ZZ"]


# --------------------------------------------------------------- validação --
def test_layout_sem_nome_da_empresa_nao_passa():
    erros = validar(_layout(socio_nome="A", socio_cpf="B"))
    assert any("Empresa" in erro for erro in erros)


def test_dados_de_socio_sem_a_coluna_do_nome_nao_passa():
    """Sem o nome do sócio toda linha seria descartada por estar incompleta —
    em silêncio, que é o pior jeito de falhar."""
    erros = validar(_layout(empresa_nome="A", socio_cpf="B", percentual_capital="C"))
    assert any("Sócio" in erro for erro in erros)


def test_mesma_coluna_em_dois_campos_nao_passa():
    erros = validar(_layout(empresa_nome="A", socio_nome="B", socio_cpf="B"))
    assert any("coluna B" in erro for erro in erros)


def test_layout_minimo_e_valido():
    assert validar(_layout(empresa_nome="A", socio_nome="B")) == []


def test_so_empresas_e_valido_sem_socio():
    assert validar(_layout(empresa_nome="A", cnpj="B", capital_social="C")) == []


def test_validacao_junta_todos_os_problemas():
    layout = LayoutImportacao("", {"socio_cpf": "9", "cnpj": "B", "capital_social": "B"}, 0)
    erros = validar(layout)
    assert len(erros) >= 5  # nome, linha, letra inválida, empresa, sócio, coluna repetida


# --------------------------------------------------------------- importação --
def test_le_colunas_fora_de_ordem_com_buracos(tmp_path):
    caminho = _planilha(tmp_path, [
        ["Cod", "Razão social", "lixo", "Nome", "Documento", "", "Perc", "Desde"],
        ["91", "Endogastro Ltda", "x", "Andre Franzotti", "076.925.727-55", "", "46,94", "23/03/2007"],
    ])
    layout = _layout(numero_chamada="A", empresa_nome="B", socio_nome="D",
                     socio_cpf="E", percentual_capital="G", data_entrada="H")

    (linha,) = importar_com_layout(caminho, layout)

    assert linha["numero_chamada"] == "91"
    assert linha["empresa_nome"] == "ENDOGASTRO LTDA"
    assert linha["socio_nome"] == "ANDRE FRANZOTTI"
    assert linha["socio_cpf"] == "076.925.727-55"
    assert linha["percentual_capital"] == 46.94
    assert linha["data_entrada"] == "2007-03-23"
    # Coluna que o layout não mapeia não é inventada: entra zerada/vazia.
    assert linha["cnpj"] == ""
    assert linha["capital_social"] == 0.0


def test_linha_inicial_pula_cabecalho_de_varias_linhas(tmp_path):
    caminho = _planilha(tmp_path, [
        ["RELATÓRIO DE SÓCIOS", None],
        ["Emitido em 20/05/2026", None],
        ["Empresa", "Sócio"],
        ["Endogastro Ltda", "Andre Franzotti"],
        ["Endogastro Ltda", "Luiza Dias"],
    ])
    layout = _layout(empresa_nome="A", socio_nome="B", linha_inicial=4)

    linhas = importar_com_layout(caminho, layout)

    assert [l["socio_nome"] for l in linhas] == ["ANDRE FRANZOTTI", "LUIZA DIAS"]


def test_linha_em_branco_antes_do_inicio_nao_desloca_a_leitura(tmp_path):
    """A pessoa conta as linhas como o Excel mostra; sumir com uma linha vazia
    faria a leitura começar no lugar errado."""
    caminho = _planilha(tmp_path, [
        [None, None],
        ["Empresa", "Sócio"],
        ["Endogastro Ltda", "Andre Franzotti"],
    ])
    layout = _layout(empresa_nome="A", socio_nome="B", linha_inicial=3)

    (linha,) = importar_com_layout(caminho, layout)

    assert linha["socio_nome"] == "ANDRE FRANZOTTI"


def test_linha_sem_empresa_e_ignorada_sem_quebrar(tmp_path):
    caminho = _planilha(tmp_path, [
        ["Empresa", "Sócio"],
        ["Endogastro Ltda", "Andre Franzotti"],
        [None, None],
        ["", "Sócio órfão"],
        ["Endogastro Ltda", "Luiza Dias"],
    ])
    linhas = importar_com_layout(caminho, _layout(empresa_nome="A", socio_nome="B"))

    assert [l["socio_nome"] for l in linhas] == ["ANDRE FRANZOTTI", "LUIZA DIAS"]


def test_erro_de_valor_aponta_a_linha_que_a_pessoa_ve(tmp_path):
    caminho = _planilha(tmp_path, [
        ["cabeçalho", "", ""],
        ["Endogastro Ltda", "Andre", "46,94"],
        ["Endogastro Ltda", "Luiza", "não é número"],
    ])
    layout = _layout(empresa_nome="A", socio_nome="B", percentual_capital="C")

    with pytest.raises(ValueError, match="Linha 3"):
        importar_com_layout(caminho, layout)


def test_importar_com_layout_invalido_nao_le_o_arquivo(tmp_path):
    caminho = _planilha(tmp_path, [["Empresa"], ["Endogastro Ltda"]])
    with pytest.raises(LayoutInvalido):
        importar_com_layout(caminho, _layout(socio_nome="A"))


def test_csv_tambem_e_lido_por_posicao(tmp_path):
    caminho = tmp_path / "origem.csv"
    caminho.write_text("Empresa;Sócio\nEndogastro Ltda;Andre Franzotti\n", encoding="utf-8")

    (linha,) = importar_com_layout(caminho, _layout(empresa_nome="A", socio_nome="B"))

    assert linha["empresa_nome"] == "ENDOGASTRO LTDA"


# ---------------------------------------------------------------- exportação --
def test_exporta_e_reimporta_pelo_mesmo_layout(tmp_path):
    """O ciclo que justifica o recurso: o arquivo sai no desenho da origem e
    volta por ele, sem ninguém remontar coluna nenhuma."""
    layout = _layout(numero_chamada="A", empresa_nome="C", socio_nome="E", socio_cpf="F",
                     percentual_capital="H", data_entrada="I", linha_inicial=3)
    caminho = tmp_path / "saida.xlsx"
    exportar_com_layout(caminho, [
        {"numero_chamada": "91", "empresa_nome": "ENDOGASTRO LTDA", "socio_nome": "ANDRE",
         "socio_cpf": "076.925.727-55", "percentual_capital": 46.94, "data_entrada": "2007-03-23"},
    ], layout)

    aba = openpyxl.load_workbook(caminho).active
    assert aba["C2"].value == "Empresa"          # cabeçalho na linha anterior à inicial
    assert aba["C3"].value == "ENDOGASTRO LTDA"  # dado na linha configurada
    assert aba["B3"].value is None               # coluna não mapeada fica vazia

    (linha,) = importar_com_layout(caminho, layout)
    assert linha["socio_cpf"] == "076.925.727-55"
    assert linha["percentual_capital"] == 46.94


def test_layout_que_comeca_na_linha_1_sai_sem_cabecalho(tmp_path):
    layout = _layout(empresa_nome="A", socio_nome="B", linha_inicial=1)
    caminho = tmp_path / "saida.xlsx"
    exportar_com_layout(caminho, [{"empresa_nome": "ENDOGASTRO LTDA", "socio_nome": "ANDRE"}], layout)

    aba = openpyxl.load_workbook(caminho).active
    assert aba["A1"].value == "ENDOGASTRO LTDA"


# -------------------------------------------------------------------- apoio --
def test_previa_mostra_so_os_campos_mapeados(tmp_path):
    caminho = _planilha(tmp_path, [
        ["Empresa", "Sócio"],
        ["Endogastro Ltda", "Andre Franzotti"],
        ["Endogastro Ltda", "Luiza Dias"],
    ])
    layout = _layout(empresa_nome="A", socio_nome="B")

    assert previa(caminho, layout) == [
        ["ENDOGASTRO LTDA", "ANDRE FRANZOTTI"],
        ["ENDOGASTRO LTDA", "LUIZA DIAS"],
    ]


def test_layout_a_partir_de_um_modelo_do_sistema():
    layout = layout_do_modelo(modelo_cadastro("empresas_socios"))

    assert layout.colunas["numero_chamada"] == "A"
    assert layout.colunas["empresa_nome"] == "B"
    assert validar(layout) == []


def test_todo_campo_do_cadastro_esta_em_algum_grupo():
    """Campo novo no cadastro sem lugar na tela sairia silenciosamente da
    importação configurável."""
    nos_grupos = [campo for grupo in GRUPOS_CAMPOS for campo in grupo.campos]
    assert sorted(nos_grupos) == sorted(CAMPOS_CADASTRO)
    assert len(nos_grupos) == len(set(nos_grupos))
    assert set(CAMPOS_POR_GRUPO) == {g.nome for g in GRUPOS_CAMPOS}
