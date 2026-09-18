"""A tela de importação/exportação com layout configurável.

O que importa aqui é o percurso que a pessoa faz: escolher o formato,
descrever as colunas da planilha dela, salvar, conferir e importar — e que o
formato escolhido valha para os dois lados, exportação inclusive.
"""
import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import sqlite3

import openpyxl
import pytest
from PySide6.QtWidgets import QApplication, QDialog, QMessageBox

from controle_lucros import db, repositories as repo
from controle_lucros.layout_importacao import LayoutImportacao
from controle_lucros.models import Empresa, Socio, VinculoSocietario
from controle_lucros.ui import importacao_cadastro_view as vista


@pytest.fixture(scope="module", autouse=True)
def app():
    aplicativo = QApplication.instance() or QApplication([])
    yield aplicativo


@pytest.fixture()
def conn():
    connection = sqlite3.connect(":memory:")
    connection.row_factory = sqlite3.Row
    db.init_schema(connection)
    yield connection
    connection.close()


@pytest.fixture()
def tela(conn):
    return vista.ImportacaoCadastroView(conn)


def _preencher(tela, nome, linha_inicial=2, **colunas):
    tela._novo_layout()
    tela.editor.nome.setText(nome)
    tela.editor.linha_inicial.setValue(linha_inicial)
    for campo, letra in colunas.items():
        tela.editor._campos[campo].setText(letra)


def _planilha(tmp_path, linhas, nome="origem.xlsx"):
    caminho = tmp_path / nome
    workbook = openpyxl.Workbook()
    aba = workbook.active
    for linha in linhas:
        aba.append(linha)
    workbook.save(caminho)
    return caminho


def _cadastro_de_exemplo(conn):
    empresa_id = repo.salvar_empresa(conn, Empresa(
        id=None, numero_chamada="91", nome="ENDOGASTRO LTDA", cnpj="",
        capital_social=1000, quantidade_cotas=100))
    socio_id = repo.salvar_socio(conn, Socio(id=None, nome="ANDRE FRANZOTTI", cpf="076.925.727-55"))
    repo.salvar_vinculo(conn, VinculoSocietario(
        id=None, empresa_id=empresa_id, socio_id=socio_id, percentual_capital=46.94,
        quantidade_cotas=46, data_entrada="2007-03-23", data_saida=None))
    return empresa_id, socio_id


# ------------------------------------------------------------------- estado --
def test_comeca_num_modelo_do_sistema_sem_grade_de_colunas(tela):
    assert tela._formato_atual()[0] == "modelo"
    assert not tela.card_layout.isVisibleTo(tela)


def test_grade_aparece_ao_criar_um_layout(tela):
    tela._novo_layout()

    assert tela._modo_layout()
    assert tela.card_layout.isVisibleTo(tela)
    assert "nome ao layout" in tela.aviso_layout.text()


def test_layout_salvo_entra_na_lista_de_formatos(tela, conn):
    _preencher(tela, "Sistema antigo", empresa_nome="B", socio_nome="D")
    tela._salvar_layout()

    assert tela.formato.currentText() == "Layout · Sistema antigo"
    assert [l.nome for l in repo.listar_layouts_importacao(conn)] == ["Sistema antigo"]
    assert "Salvo" in tela.aviso_layout.text()


def test_salvar_recolhe_a_grade_de_colunas(tela):
    """Configurado o layout, a tela volta a ser só escolher o formato e
    apontar o arquivo — a grade não fica aberta embaixo atrapalhando."""
    _preencher(tela, "Sistema antigo", empresa_nome="B", socio_nome="D")
    tela._salvar_layout()

    assert not tela.card_layout.isVisibleTo(tela)
    assert tela.btn_editar_layout.isEnabled()
    assert "Sistema antigo" in tela.resultado.text()


def test_editar_layout_reabre_a_grade_e_fechar_recolhe_de_novo(tela):
    _preencher(tela, "Sistema antigo", empresa_nome="B", socio_nome="D")
    tela._salvar_layout()

    tela._editar_layout()
    assert tela.card_layout.isVisibleTo(tela)
    assert not tela.btn_editar_layout.isEnabled()  # já está aberta

    tela._fechar_editor()
    assert not tela.card_layout.isVisibleTo(tela)
    # Recolher não desfaz nada: o layout escolhido continua valendo.
    assert tela._layout_escolhido().colunas == {"empresa_nome": "B", "socio_nome": "D"}


def test_modelo_do_sistema_nao_oferece_editar_layout(tela):
    assert tela._formato_atual()[0] == "modelo"
    assert not tela.btn_editar_layout.isEnabled()


def test_layout_incompleto_avisa_e_nao_salva(tela, conn, monkeypatch):
    avisos = []
    monkeypatch.setattr(vista.QMessageBox, "warning",
                        staticmethod(lambda parent, titulo, texto, *a: avisos.append(texto)))
    _preencher(tela, "Faltando empresa", socio_nome="D")
    tela._salvar_layout()

    assert avisos and "Empresa" in avisos[0]
    assert repo.listar_layouts_importacao(conn) == []


def test_editar_layout_salvo_avisa_que_ha_mudanca_pendente(tela):
    _preencher(tela, "Sistema antigo", empresa_nome="B", socio_nome="D")
    tela._salvar_layout()

    tela._editar_layout()
    tela.editor._campos["cnpj"].setText("F")

    assert "não salvas" in tela.aviso_layout.text()


def test_duplicar_modelo_do_sistema_gera_layout_preenchido(tela):
    tela.formato.setCurrentIndex(tela.formato.findData("modelo:empresas"))
    tela._duplicar_como_layout()

    layout = tela._layout_escolhido()
    assert layout.id is None
    assert layout.colunas["numero_chamada"] == "A"
    assert layout.colunas["empresa_nome"] == "B"
    assert tela.card_layout.isVisibleTo(tela)


def test_excluir_layout_volta_para_o_modelo_padrao(tela, conn, monkeypatch):
    _preencher(tela, "Sistema antigo", empresa_nome="B", socio_nome="D")
    tela._salvar_layout()
    monkeypatch.setattr(vista.QMessageBox, "question", staticmethod(lambda *a, **k: QMessageBox.Yes))

    tela._excluir_layout()

    assert repo.listar_layouts_importacao(conn) == []
    assert tela._formato_atual()[0] == "modelo"
    assert not tela.card_layout.isVisibleTo(tela)


# ---------------------------------------------------------------- importar --
def test_importa_planilha_de_outra_origem_pelo_layout(tela, conn, tmp_path, monkeypatch):
    """O caso que motiva o recurso: planilha com as colunas em outra ordem,
    cabeçalho de duas linhas e uma coluna que o sistema nem usa."""
    caminho = _planilha(tmp_path, [
        ["RELATÓRIO DE SÓCIOS — SISTEMA ANTIGO", None, None, None, None],
        ["Cod", "Razão social", "Observações", "Nome do sócio", "CPF"],
        ["91", "Endogastro Ltda", "qualquer coisa", "Andre Franzotti", "076.925.727-55"],
        ["91", "Endogastro Ltda", "", "Luiza Dias", "103.285.827-35"],
    ])
    _preencher(tela, "Sistema antigo", linha_inicial=3,
               numero_chamada="A", empresa_nome="B", socio_nome="D", socio_cpf="E")
    tela._salvar_layout()

    monkeypatch.setattr(vista.QFileDialog, "getOpenFileName",
                        staticmethod(lambda *a, **k: (str(caminho), "")))
    monkeypatch.setattr(vista.QMessageBox, "information", staticmethod(lambda *a, **k: None))
    monkeypatch.setattr(vista._DialogoPrevia, "exec", lambda self: QDialog.Accepted)
    monkeypatch.setattr(vista.DialogoRevisaoCadastro, "exec",
                        lambda self: (self._cadastrar_todos(), QDialog.Accepted)[1])
    monkeypatch.setattr(vista.QMessageBox, "question", staticmethod(lambda *a, **k: QMessageBox.Yes))

    tela._importar_planilha()

    (empresa,) = repo.listar_empresas(conn)
    assert empresa.nome == "ENDOGASTRO LTDA"
    assert empresa.numero_chamada == "91"
    assert {s.nome for s in repo.listar_socios(conn)} == {"ANDRE FRANZOTTI", "LUIZA DIAS"}
    assert len(repo.listar_vinculos_empresa(conn, empresa.id)) == 2


def test_recusar_a_previa_nao_grava_nada(tela, conn, tmp_path, monkeypatch):
    caminho = _planilha(tmp_path, [["Empresa", "Sócio"], ["Endogastro Ltda", "Andre"]])
    _preencher(tela, "Simples", empresa_nome="A", socio_nome="B")
    monkeypatch.setattr(vista.QFileDialog, "getOpenFileName",
                        staticmethod(lambda *a, **k: (str(caminho), "")))
    monkeypatch.setattr(vista._DialogoPrevia, "exec", lambda self: QDialog.Rejected)

    tela._importar_planilha()

    assert repo.listar_empresas(conn) == []


def test_previa_mostra_a_leitura_sem_importar(tela, conn, tmp_path, monkeypatch):
    caminho = _planilha(tmp_path, [["Empresa", "Sócio"], ["Endogastro Ltda", "Andre"]])
    _preencher(tela, "Simples", empresa_nome="A", socio_nome="B")
    monkeypatch.setattr(vista.QFileDialog, "getOpenFileName",
                        staticmethod(lambda *a, **k: (str(caminho), "")))

    vistos = {}

    def abrir(self):
        vistos["titulo"] = self.windowTitle()
        return QDialog.Rejected

    monkeypatch.setattr(vista._DialogoPrevia, "exec", abrir)
    tela._conferir_layout()

    assert vistos["titulo"] == "Conferir leitura da planilha"
    assert repo.listar_empresas(conn) == []


def test_erro_de_leitura_avisa_em_vez_de_quebrar(tela, tmp_path, monkeypatch):
    caminho = _planilha(tmp_path, [["Empresa", "%"], ["Endogastro Ltda", "isto não é número"]])
    _preencher(tela, "Simples", empresa_nome="A", socio_nome="A", percentual_capital="B")
    avisos = []
    monkeypatch.setattr(vista.QFileDialog, "getOpenFileName",
                        staticmethod(lambda *a, **k: (str(caminho), "")))
    monkeypatch.setattr(vista.QMessageBox, "warning",
                        staticmethod(lambda parent, titulo, texto, *a: avisos.append(texto)))

    tela._conferir_layout()

    assert avisos  # o layout tem a mesma coluna em dois campos: reclama antes de ler


# ---------------------------------------------------------------- exportar --
def test_exporta_cadastro_no_desenho_do_layout(tela, conn, tmp_path, monkeypatch):
    _cadastro_de_exemplo(conn)
    destino = tmp_path / "saida.xlsx"
    _preencher(tela, "Sistema antigo", linha_inicial=3,
               numero_chamada="A", empresa_nome="C", socio_nome="E", socio_cpf="F",
               percentual_capital="H", data_entrada="I")
    monkeypatch.setattr(vista.QFileDialog, "getSaveFileName",
                        staticmethod(lambda *a, **k: (str(destino), "")))
    monkeypatch.setattr(vista.QMessageBox, "information", staticmethod(lambda *a, **k: None))

    tela._exportar_atual()

    aba = openpyxl.load_workbook(destino).active
    assert aba["C2"].value == "Empresa"
    assert aba["C3"].value == "ENDOGASTRO LTDA"
    assert aba["F3"].value == "076.925.727-55"
    assert aba["B3"].value is None


def test_exportar_em_branco_no_layout_sai_so_com_cabecalho(tela, tmp_path, monkeypatch):
    destino = tmp_path / "branco.xlsx"
    _preencher(tela, "Sistema antigo", empresa_nome="C", socio_nome="E")
    monkeypatch.setattr(vista.QFileDialog, "getSaveFileName",
                        staticmethod(lambda *a, **k: (str(destino), "")))
    monkeypatch.setattr(vista.QMessageBox, "information", staticmethod(lambda *a, **k: None))

    tela._exportar_modelo()

    aba = openpyxl.load_workbook(destino).active
    assert aba["C1"].value == "Empresa"
    assert aba["C2"].value is None


def test_modelo_do_sistema_continua_exportando_pelo_cabecalho(tela, conn, tmp_path, monkeypatch):
    _cadastro_de_exemplo(conn)
    destino = tmp_path / "modelo.xlsx"
    tela.formato.setCurrentIndex(tela.formato.findData("modelo:empresas_socios"))
    monkeypatch.setattr(vista.QFileDialog, "getSaveFileName",
                        staticmethod(lambda *a, **k: (str(destino), "")))
    monkeypatch.setattr(vista.QMessageBox, "information", staticmethod(lambda *a, **k: None))

    tela._exportar_atual()

    aba = openpyxl.load_workbook(destino).active
    assert aba["A1"].value == "Nº Empresa"
    assert aba["B2"].value == "ENDOGASTRO LTDA"
    # O modelo do sistema traz a aba de exemplo; o layout, não.
    assert "Exemplo" in openpyxl.load_workbook(destino).sheetnames


def test_ciclo_completo_exporta_reimporta_e_nao_duplica(tela, conn, tmp_path, monkeypatch):
    """Exportar e reimportar pelo mesmo layout não pode criar nada de novo —
    é a regra de "o que já existe não é importado de novo"."""
    empresa_id, _socio_id = _cadastro_de_exemplo(conn)
    destino = tmp_path / "ciclo.xlsx"
    _preencher(tela, "Ciclo", numero_chamada="A", empresa_nome="B", socio_nome="C",
               socio_cpf="D", percentual_capital="E", data_entrada="F")
    monkeypatch.setattr(vista.QFileDialog, "getSaveFileName",
                        staticmethod(lambda *a, **k: (str(destino), "")))
    monkeypatch.setattr(vista.QFileDialog, "getOpenFileName",
                        staticmethod(lambda *a, **k: (str(destino), "")))
    monkeypatch.setattr(vista.QMessageBox, "information", staticmethod(lambda *a, **k: None))
    monkeypatch.setattr(vista._DialogoPrevia, "exec", lambda self: QDialog.Accepted)

    tela._exportar_atual()
    tela._importar_planilha()

    assert len(repo.listar_empresas(conn)) == 1
    assert len(repo.listar_socios(conn)) == 1
    assert len(repo.listar_vinculos_empresa(conn, empresa_id)) == 1
    assert "1 já existiam (ignorados)" in tela.resultado.text()


# ---------------------------------------------------- reaproveitar o layout --
def test_layout_salvo_aparece_ao_abrir_a_tela_de_novo(conn):
    """O uso real é este: configurou uma vez, e todo mês só escolhe o formato
    e aponta o arquivo."""
    repo.salvar_layout_importacao(
        conn, LayoutImportacao("Sistema antigo", {"empresa_nome": "B", "socio_nome": "D"}, 3)
    )
    tela = vista.ImportacaoCadastroView(conn)

    tela.formato.setCurrentIndex(tela.formato.findData("layout:1"))

    assert tela.formato.currentText() == "Layout · Sistema antigo"
    # A grade vem recolhida: escolher o formato já basta pra importar por ele.
    assert not tela.card_layout.isVisibleTo(tela)
    assert "Editar layout" in tela.descricao_formato.text()
    assert tela.editor.layout_atual().colunas == {"empresa_nome": "B", "socio_nome": "D"}
    assert tela.editor.linha_inicial.value() == 3
    assert "Salvo" in tela.aviso_layout.text()


def test_tela_reabre_no_ultimo_formato_usado(conn):
    """Voltar sempre no modelo completo faria quem importa da mesma origem
    reencontrar o layout toda vez — e importar com o formato errado é
    silencioso. (A preferência vai para a pasta isolada do conftest.)"""
    repo.salvar_layout_importacao(conn, LayoutImportacao("Sistema antigo", {"empresa_nome": "B", "socio_nome": "D"}))

    primeira = vista.ImportacaoCadastroView(conn)
    primeira.formato.setCurrentIndex(primeira.formato.findData("layout:1"))

    segunda = vista.ImportacaoCadastroView(conn)
    assert segunda.formato.currentText() == "Layout · Sistema antigo"
    assert segunda.editor.layout_atual().colunas == {"empresa_nome": "B", "socio_nome": "D"}


def test_atualizar_repoe_a_lista_sem_perder_a_selecao(conn):
    repo.salvar_layout_importacao(conn, LayoutImportacao("Sistema antigo", {"empresa_nome": "B", "socio_nome": "D"}))
    tela = vista.ImportacaoCadastroView(conn)
    tela.formato.setCurrentIndex(tela.formato.findData("layout:1"))

    repo.salvar_layout_importacao(conn, LayoutImportacao("Contabilidade X", {"empresa_nome": "A", "socio_nome": "B"}))
    tela.atualizar()

    textos = [tela.formato.itemText(i) for i in range(tela.formato.count())]
    assert "Layout · Contabilidade X" in textos
    assert tela.formato.currentText() == "Layout · Sistema antigo"
